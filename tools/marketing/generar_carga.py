"""Genera carga_historial.sql desde el Calendario Maestro.
Uso: python3 -I generar_carga.py <ruta.xlsx> <carpeta_salida>
Reemplazar __ORG__ en el SQL por el id de la organización. No se versiona el .xlsx."""
import openpyxl, json, sys, datetime as dt
from collections import defaultdict
import sys
F=sys.argv[1]
wb=openpyxl.load_workbook(F,data_only=True)
ig=[]; ws=wb["Datos Metricool"]
for r in ws.iter_rows(min_row=4,values_only=True):
    if isinstance(r[0],dt.datetime): ig.append(dict(red="instagram",fecha=r[0].date(),pieza=r[1],pilar=r[2],formato=r[3],valor=r[4],interacciones=r[5]))
fb=[];li=[]; ws=wb["Datos FB y LinkedIn"]
for r in ws.iter_rows(min_row=5,values_only=True):
    if isinstance(r[0],dt.datetime): fb.append(dict(red="facebook",fecha=r[0].date(),pieza=r[1],pilar=r[2],valor=r[3],reacciones=r[4],compartidos=r[5],clics=r[6]))
    if isinstance(r[8],dt.datetime): li.append(dict(red="linkedin",fecha=r[8].date(),pieza=r[9],pilar=r[10],valor=r[11],reacciones=r[12],compartidos=r[13],clics=r[14]))
FALLIDAS=[dt.date(2026,9,d) for d in (22,23,24,26,29)]

import json, statistics
# Agrupar por (fecha, tema): 28 ago tiene dos publicaciones de LinkedIn; "Día de los Abuelos" no tiene contraparte en FB.
piezas=[]
by=defaultdict(list)
for x in ig+fb+li: by[x['fecha']].append(x)
for f in sorted(by):
    xs=by[f]
    li_x=[x for x in xs if x['red']=='linkedin']
    otros=[x for x in xs if x['red']!='linkedin']
    if len(li_x)==2:
        fbx=[x for x in otros if x['red']=='facebook'][0]
        match=[x for x in li_x if x['pieza']==fbx['pieza']]
        assert match, (f,)
        grupos=[otros+match, [x for x in li_x if x not in match]]
    else: grupos=[xs]
    for g in grupos: piezas.append((f,g))
rows=[]
for i,(f,g) in enumerate(piezas):
    d={x['red']:x for x in g}
    ig_,fb_,li_=d.get('instagram'),d.get('facebook'),d.get('linkedin')
    pil={x['pilar'] for x in g}; assert len(pil)==1,(f,pil)
    titulo=(ig_ or fb_ or li_)['pieza']
    cod=f"H-{f:%Y%m%d}"+("-2" if sum(1 for q in piezas[:i] if q[0]==f)>0 else "")
    ig_estado = 'E' if (f in FALLIDAS) else None
    rows.append([cod,str(f),titulo,pil.pop(),(ig_ or {}).get('formato'),
        ig_['valor'] if ig_ else ig_estado, ig_['interacciones'] if ig_ else None,
        fb_['valor'] if fb_ else None, fb_['reacciones'] if fb_ else None, fb_['compartidos'] if fb_ else None, fb_['clics'] if fb_ else None,
        li_['valor'] if li_ else None, li_['reacciones'] if li_ else None, li_['compartidos'] if li_ else None, li_['clics'] if li_ else None])
json.dump(rows,open(sys.argv[2]+'/piezas.json','w'),ensure_ascii=False,separators=(',',':'))
print("piezas:",len(rows))
print("IG pub:",sum(1 for r in rows if isinstance(r[5],(int,float))),"IG error:",sum(1 for r in rows if r[5]=='E'),"FB:",sum(1 for r in rows if r[7] is not None),"LI:",sum(1 for r in rows if r[11] is not None))
v=[r[5] for r in rows if isinstance(r[5],(int,float))]
print("IG media",statistics.mean(v),"mediana",statistics.median(v))
print("FB media",statistics.mean(r[7] for r in rows if r[7] is not None),"LI media",statistics.mean(r[11] for r in rows if r[11] is not None))
print(len(json.dumps(rows,ensure_ascii=False,separators=(',',':'))),"bytes")

import json
# Cifras semanales de Instagram que trae el archivo (Datos y decisiones, sección 2)
ws=wb["Datos y decisiones"]; ig_sem=[]
on=False
for r in ws.iter_rows(values_only=True):
    if r[1]=="Semana (lunes)": on=True; continue
    if on:
        if isinstance(r[1],dt.datetime): ig_sem.append([str(r[1].date()), r[2], None if isinstance(r[3],str) else r[3]])
        else: break
print(ig_sem)
P=json.dumps(rows,ensure_ascii=False,separators=(',',':')).replace("'","''")
S=json.dumps(ig_sem,ensure_ascii=False,separators=(',',':'))
sql=r"""DO $load$
DECLARE
  v_org uuid := '__ORG__';
  r jsonb; v_pid uuid; v_prid uuid; v_f date; v_ts timestamptz;
  v_caso uuid;
  w date; v_red text; s record; v_ceros integer; v_prov boolean; v_nota text; z record; v_caso_v uuid;
BEGIN
  PERFORM public.mk_sembrar_reglas_v1(v_org);

  FOR r IN SELECT jsonb_array_elements($j$__PIEZAS__$j$::jsonb) LOOP
    v_f := (r->>1)::date;
    v_ts := (v_f::text || ' 11:00')::timestamp AT TIME ZONE 'America/Mexico_City';
    INSERT INTO public.mk_pieza (organization_id, codigo, titulo, pilar, formato, etapa, fecha, origen)
    VALUES (v_org, r->>0, r->>2, r->>3, NULLIF(r->>4, ''), 1, v_f, 'calendario_maestro')
    RETURNING id INTO v_pid;

    -- Instagram: valor numérico = publicada; 'E' = con errores (error «Unknown» según la pestaña Monitoreo)
    IF r->5 IS NOT NULL AND jsonb_typeof(r->5) = 'number' THEN
      INSERT INTO public.mk_pieza_red (organization_id, pieza_id, red, estado, publicada_at, hora_estimada)
      VALUES (v_org, v_pid, 'instagram', 'publicada', v_ts, true) RETURNING id INTO v_prid;
      INSERT INTO public.mk_resultado (organization_id, pieza_red_id, corte, fuente, medida, valor, extra, capturado_at)
      VALUES (v_org, v_prid, 'acumulado', 'calendario_maestro', 'alcance', (r->>5)::numeric,
              jsonb_build_object('interacciones', (r->>6)::int), '2026-10-31 12:00 America/Mexico_City');
    ELSIF r->>5 = 'E' THEN
      INSERT INTO public.mk_pieza_red (organization_id, pieza_id, red, estado, mensaje_error, hora_estimada)
      VALUES (v_org, v_pid, 'instagram', 'con_errores', 'Unknown', true);
    END IF;

    IF r->7 IS NOT NULL AND jsonb_typeof(r->7) = 'number' THEN
      INSERT INTO public.mk_pieza_red (organization_id, pieza_id, red, estado, publicada_at, hora_estimada)
      VALUES (v_org, v_pid, 'facebook', 'publicada', v_ts, true) RETURNING id INTO v_prid;
      INSERT INTO public.mk_resultado (organization_id, pieza_red_id, corte, fuente, medida, valor, reacciones, compartidos, clics, comentarios, capturado_at)
      VALUES (v_org, v_prid, 'acumulado', 'calendario_maestro', 'alcance', (r->>7)::numeric,
              (r->>8)::int, (r->>9)::int, (r->>10)::int, 0, '2026-10-31 12:00 America/Mexico_City');
    END IF;

    IF r->11 IS NOT NULL AND jsonb_typeof(r->11) = 'number' THEN
      INSERT INTO public.mk_pieza_red (organization_id, pieza_id, red, estado, publicada_at, hora_estimada)
      VALUES (v_org, v_pid, 'linkedin', 'publicada', v_ts, true) RETURNING id INTO v_prid;
      INSERT INTO public.mk_resultado (organization_id, pieza_red_id, corte, fuente, medida, valor, reacciones, compartidos, clics, comentarios, capturado_at)
      VALUES (v_org, v_prid, 'acumulado', 'calendario_maestro', 'impresiones', (r->>11)::numeric,
              (r->>12)::int, (r->>13)::int, (r->>14)::int, 0, '2026-10-31 12:00 America/Mexico_City');
    END IF;
  END LOOP;

  -- Seguidores: una sola lectura (consulta del 31 oct 2026), puesta en la semana del 26 oct
  INSERT INTO public.mk_seguidores_semana (organization_id, red, semana_inicio, seguidores, fuente) VALUES
    (v_org, 'instagram', '2026-10-26', 154, 'calendario_maestro'),
    (v_org, 'facebook',  '2026-10-26', 1111, 'calendario_maestro'),
    (v_org, 'linkedin',  '2026-10-26', 1295, 'calendario_maestro');

  -- Cobertura de la fuente (qué periodo trae el archivo por red). Instagram: la hoja empieza el 8 sep.
  INSERT INTO public.mk_cobertura_fuente (organization_id, red, fuente, desde, hasta, nota) VALUES
    (v_org, 'instagram', 'calendario_maestro', '2026-09-08', '2026-10-29', 'Hoja Datos Metricool: 26 piezas; no trae 30 jul – 7 sep ni las 5 piezas con error'),
    (v_org, 'facebook',  'calendario_maestro', '2026-07-30', '2026-10-30', 'Hoja Datos FB y LinkedIn: 47 posts'),
    (v_org, 'linkedin',  'calendario_maestro', '2026-07-30', '2026-10-30', 'Hoja Datos FB y LinkedIn: 48 posts');

  -- Caso 1 (septiembre)
  INSERT INTO public.mk_caso (organization_id, tipo, red, titulo, descripcion, evidencia, estado, responsable, detectado_at)
  VALUES (v_org, 'investigacion', 'instagram',
    'Septiembre 2026: 5 piezas fallidas en Instagram, 11 días sin posts y caída de alcance desde el 17 de septiembre',
    'Caso cargado desde el Calendario Maestro (pestañas Monitoreo, Datos Metricool, Datos FB y LinkedIn y Datos y decisiones; consulta de Metricool del 31 oct 2026). La caída de alcance empezó antes de las fallas; la causa no está determinada. El Revisor lo recorre en la Entrega 4 sin ejecutar correcciones.',
    $e$[
{"tipo":"hecho","dato":"Las 5 piezas del 22, 23, 24, 26 y 29 de septiembre figuran «Con errores» en Instagram, mensaje «Unknown»; en Facebook y LinkedIn sí se publicaron.","fuente":"Calendario Maestro · Monitoreo, sección 1; Datos Metricool, nota"},
{"tipo":"hecho","dato":"Instagram estuvo 11 días sin posts en el feed (del 19 al 30 de septiembre) y no se detectó a tiempo.","fuente":"Calendario Maestro · Monitoreo, sección 1 y 4"},
{"tipo":"hecho","dato":"Alcance en Instagram: 137 el 16 sep, 18 el 17 sep, 14 el 19 sep; casi todo octubre entre 7 y 29 (con picos de 129 el 13 oct y 79 el 17 oct).","fuente":"Calendario Maestro · Datos Metricool"},
{"tipo":"hecho","dato":"Semana del 21 sep sin datos de Instagram: las piezas fallaron. Promedio semanal de IG: 61 (7 sep), 58.5 (14 sep), sin datos (21 sep), 17.67 (28 sep), 25.75 (5 oct), 57.5 (12 oct), 10.25 (19 oct), 8 (26 oct).","fuente":"Calendario Maestro · Datos y decisiones, sección 2"},
{"tipo":"hecho","dato":"El mensaje de error es solo «Unknown»: no se sabe si fue la conexión, Instagram o las piezas.","fuente":"Calendario Maestro · Monitoreo, sección 1"},
{"tipo":"hecho","dato":"La caída del 17 y 19 de septiembre ocurrió antes de las fallas, por lo que la falla técnica no la explica por completo.","fuente":"Calendario Maestro · Monitoreo, sección 1"},
{"tipo":"hecho","dato":"Todas las piezas salieron cerca de las 11:00 (una sola hora); Instagram tiene 154 seguidores.","fuente":"Calendario Maestro · Datos FB y LinkedIn, nota; Datos y decisiones"},
{"tipo":"hipotesis","estado":"Confirmada, parcial","dato":"1. Falla técnica: la pieza no se publicó. Confirmada para 5 piezas de IG.","fuente":"Calendario Maestro · Monitoreo, sección 4"},
{"tipo":"hipotesis","estado":"Por probar","dato":"2. Hueco de 11 días en Instagram (19 → 30 sep). Plausible, no probada.","fuente":"Calendario Maestro · Monitoreo, sección 4"},
{"tipo":"hipotesis","estado":"Probable","dato":"3. El 16 de septiembre fue un pico por la fecha, no el nivel normal (mediana 15).","fuente":"Calendario Maestro · Monitoreo, sección 4"},
{"tipo":"hipotesis","estado":"En prueba","dato":"4. Tema: Contable ≈14, Legal ≈18, Frase 8, frente a PLD ≈51.","fuente":"Calendario Maestro · Monitoreo, sección 4"},
{"tipo":"hipotesis","estado":"Por probar","dato":"5. Copy: 9-10 hashtags y link kawiil.mx en cada pieza. Sin dato todavía.","fuente":"Calendario Maestro · Monitoreo, sección 4"},
{"tipo":"hipotesis","estado":"Por probar","dato":"6. Diseño: carruseles con mucho texto y fondo oscuro.","fuente":"Calendario Maestro · Monitoreo, sección 4"},
{"tipo":"hipotesis","estado":"En prueba","dato":"7. Horario: todo a las 11:00; sin datos de otros horarios.","fuente":"Calendario Maestro · Monitoreo, sección 4"},
{"tipo":"hipotesis","estado":"Contexto","dato":"8. Cuenta chica: 154 seguidores en Instagram.","fuente":"Calendario Maestro · Monitoreo, sección 4"},
{"tipo":"hipotesis","estado":"Pendiente","dato":"9. Estado de la cuenta de Instagram (restricción). No se ha revisado.","fuente":"Calendario Maestro · Monitoreo, sección 4"}
]$e$::jsonb,
    'abierto', 'revisor', '2026-10-31 12:00 America/Mexico_City')
  RETURNING id INTO v_caso;

  -- Casos de verificación: pieza con 0 de alcance/impresiones y más de 48 h desde su publicación (regla dato_sospechoso).
  -- El cero cuenta en el promedio del semáforo. «Más de 48 h» se mide contra la consulta del 31 oct 12:00 (hora estimada).
  FOR z IN
    SELECT p.id AS pieza_id, p.titulo, p.fecha, r.red, res.valor, r.publicada_at
    FROM public.mk_resultado res
    JOIN public.mk_pieza_red r ON r.id = res.pieza_red_id
    JOIN public.mk_pieza p ON p.id = r.pieza_id
    WHERE p.organization_id = v_org AND res.valor = 0 AND r.estado = 'publicada'
      AND r.publicada_at < res.capturado_at - interval '48 hours'
    ORDER BY p.fecha, r.red
  LOOP
    INSERT INTO public.mk_caso (organization_id, tipo, regla_codigo, red, pieza_id, titulo, descripcion, evidencia, estado, responsable, detectado_at)
    VALUES (v_org, 'dato_sospechoso', 'dato_sospechoso', z.red, z.pieza_id,
      'Dato sospechoso, verificar si se publicó: «' || z.titulo || '» (' || z.fecha || ') con 0 en ' || z.red,
      'Dato sospechoso, verificar si se publicó. Cero ' || CASE WHEN z.red = 'linkedin' THEN 'impresiones' ELSE 'de alcance' END ||
      ' con más de 48 horas desde la publicación. El cero cuenta en el promedio del semáforo. Puede ser retraso de datos o una falla de publicación; no se trata como falsa alarma hasta verificar.',
      jsonb_build_array(
        jsonb_build_object('tipo', 'hecho', 'dato', 'Valor 0 en ' || z.red || ' para la pieza del ' || z.fecha, 'fuente', 'Calendario Maestro · Datos FB y LinkedIn'),
        jsonb_build_object('tipo', 'hecho', 'dato', 'El archivo pide confirmar si estos posts se publicaron o es retraso de datos', 'fuente', 'Calendario Maestro · Datos FB y LinkedIn, nota')
      ),
      'abierto', 'vigilante', '2026-10-31 12:00 America/Mexico_City')
    RETURNING id INTO v_caso_v;
    IF z.fecha = '2026-09-22' THEN
      UPDATE public.mk_caso SET evidencia = evidencia || jsonb_build_array(jsonb_build_object('tipo', 'hecho',
        'dato', 'El 22 de septiembre es también una de las 5 fechas en que la pieza figura «Con errores» en Instagram', 'fuente', 'Calendario Maestro · Monitoreo, sección 1'))
      WHERE id = v_caso_v;
    END IF;
  END LOOP;

  -- Calibración del semáforo: fórmula vs cifras del archivo (solo Instagram trae cifra semanal)
  FOR w IN SELECT generate_series('2026-07-27'::date, '2026-10-26'::date, interval '7 days')::date LOOP
    SELECT EXISTS (
      SELECT 1 FROM public.mk_pieza p
      WHERE p.organization_id = v_org AND p.primera_publicacion_at IS NOT NULL
        AND public.mk_semana_inicio(p.primera_publicacion_at) = w
        AND p.primera_publicacion_at + interval '7 days' > '2026-10-31 12:00 America/Mexico_City'::timestamptz
    ) INTO v_prov;
    FOREACH v_red IN ARRAY ARRAY['instagram', 'facebook', 'linkedin'] LOOP
      SELECT * INTO s FROM public.mk_semaforo_semana(v_org, w, v_red, 'calendario_maestro');
      SELECT count(*) INTO v_ceros FROM public.mk_caso c JOIN public.mk_pieza p ON p.id = c.pieza_id
      WHERE c.organization_id = v_org AND c.tipo = 'dato_sospechoso' AND c.red = v_red
        AND public.mk_semana_inicio(p.primera_publicacion_at) = w;
      v_nota := concat_ws(' ',
        CASE WHEN v_prov THEN 'Lectura provisional (hay piezas con menos de 7 días).' END,
        CASE WHEN v_ceros > 0 THEN v_ceros || ' pieza(s) con 0: dato sospechoso, verificar si se publicó.' END,
        CASE WHEN s.motivo_gris = 'sin_datos_fuente' THEN 'Gris: sin datos de la fuente; recalcular cuando llegue el historial de Metricool.'
             WHEN s.motivo_gris = 'sin_piezas_publicadas' THEN 'Gris: sin piezas publicadas (si había piezas programadas, el caso lo abren las reglas 1 y 3).' END);
      INSERT INTO public.mk_calibracion
        (organization_id, semana_inicio, red, piezas_calculadas, promedio_calculado, piezas_calendario, promedio_calendario,
         color_calculado, color_calendario, motivo_gris, diferencia, nota)
      SELECT v_org, w, v_red, s.piezas, s.promedio,
             c.piezas, c.promedio, s.color, NULL, s.motivo_gris,
             CASE WHEN c.piezas IS NOT NULL THEN s.promedio - c.promedio END,
             concat_ws(' ',
               CASE WHEN v_red <> 'instagram' THEN 'El archivo no trae cifra semanal de esta red; calibrar contra la API.'
                    WHEN c.piezas IS NULL THEN 'El archivo no trae cifra de Instagram para esta semana.' END,
               NULLIF(v_nota, ''))
      FROM (SELECT 1) x
      LEFT JOIN (SELECT (e->>1)::int AS piezas, (e->>2)::numeric AS promedio
                 FROM jsonb_array_elements($s$__SEM__$s$::jsonb) e WHERE (e->>0)::date = w AND v_red = 'instagram') c ON true;
    END LOOP;
  END LOOP;

  INSERT INTO public.mk_bitacora (organization_id, agente, accion, resultado, detalle, caso_id)
  VALUES (v_org, 'sistema', 'carga_historial_calendario_maestro',
          '48 piezas; 26+5 publicaciones IG, 47 FB, 48 LI; 121 resultados; caso 1; casos de verificación de ceros; cobertura; calibración',
          jsonb_build_object('fuente', 'Calendario_Maestro_Kawiil_Nov2026_Ene2027.xlsx', 'metricool_api', 'pendiente'), v_caso);
END
$load$;
"""
sql=sql.replace("__PIEZAS__",P).replace("__SEM__",S)
open(sys.argv[2]+'/carga_historial.sql','w').write(sql); print(len(sql))
